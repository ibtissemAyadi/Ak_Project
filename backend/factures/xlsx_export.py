"""Export Excel (.xlsx) d'une facture — même mise en forme que le PDF
(en-tête bleu marine, lignes alternées, total mis en évidence)."""

import io
from decimal import Decimal

from django.conf import settings

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from .models import MODE_REGLEMENT_CHOICES

COMPANY_NAME_CONTACT = 'A and K CONSEIL ET INGENIERIE'
COMPANY_ADDRESS = '13 rue Ali Belhouane, La Soukra, 2036 Ariana – Tunisie'
COMPANY_VAT_ID = '1931449T/A/M/000'
COMPANY_EMAIL = 'contact@ak-ingenierie.fr'
LOGO_PATH = settings.BASE_DIR / 'devis' / 'assets' / 'ak_logo.png'

BANQUE_NOM = 'UIB – Union Internationale de Banque'
BANQUE_IBAN = 'TN59 12 206 00 00055002674 77'
BANQUE_BIC = 'UIBKTNTT'

COULEUR_PRINCIPALE = '1F3B57'
COULEUR_FOND_ALT = 'F7F9FB'
COULEUR_FOND_TOTAL = 'EEF3F7'
COULEUR_BORDURE = 'D8DEE4'
COULEUR_TEXTE_ATTENUE = '5B5B5B'

MOIS_FR = [
    'janvier', 'février', 'mars', 'avril', 'mai', 'juin',
    'juillet', 'août', 'septembre', 'octobre', 'novembre', 'décembre',
]
MODE_REGLEMENT_LABELS = dict(MODE_REGLEMENT_CHOICES)


def formater_date_longue(d) -> str:
    return f'{d.day} {MOIS_FR[d.month - 1]} {d.year}' if d else '—'


FONT_TITRE = Font(name='Calibri', size=20, bold=True, color=COULEUR_PRINCIPALE)
FONT_NORMAL = Font(name='Calibri', size=10)
FONT_MUTED = Font(name='Calibri', size=10, color=COULEUR_TEXTE_ATTENUE)
FONT_ENTETE_TABLE = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
FONT_TOTAL_LABEL = Font(name='Calibri', size=12, bold=True)
FONT_TOTAL_VALEUR = Font(name='Calibri', size=12, bold=True)
FONT_MERCI = Font(name='Calibri', size=10, bold=True, color=COULEUR_PRINCIPALE)
FONT_ENTREPRISE_NOM = Font(name='Calibri', size=10, bold=True)
FONT_CHAMP_DROITE = Font(name='Calibri', size=10, bold=True, color=COULEUR_PRINCIPALE)
FONT_A_CLIENT = Font(name='Calibri', size=11, bold=True, color=COULEUR_PRINCIPALE)
FONT_LABEL = Font(name='Calibri', size=10, bold=True)

FILL_ENTETE_TABLE = PatternFill('solid', fgColor=COULEUR_PRINCIPALE)
FILL_ALT = PatternFill('solid', fgColor=COULEUR_FOND_ALT)
FILL_TOTAL = PatternFill('solid', fgColor=COULEUR_FOND_TOTAL)

BORDURE_BASSE = Border(bottom=Side(style='thin', color=COULEUR_BORDURE))
BORDURE_PRINCIPALE = Border(bottom=Side(style='thin', color=COULEUR_PRINCIPALE))
# Grille complète (lignes ET colonnes séparées) pour les cellules du tableau
# de prestations — un tableau bien structuré, pas seulement des séparateurs horizontaux.
BORDURE_GRILLE = Border(
    top=Side(style='thin', color=COULEUR_BORDURE), bottom=Side(style='thin', color=COULEUR_BORDURE),
    left=Side(style='thin', color=COULEUR_BORDURE), right=Side(style='thin', color=COULEUR_BORDURE),
)
FONT_TITRE_PROJET_LIGNE = Font(name='Calibri', size=10, bold=True, color=COULEUR_PRINCIPALE)

FORMAT_MONTANT = '#,##0.00" €"'
FORMAT_NOMBRE = '#,##0.00'

# Ordre demandé : Quantité, Description, Prix unitaire, Total.
COLONNES_TABLEAU = ['Quantité', 'Description', 'Prix unitaire', 'Total']
LARGEURS_COLONNES = [12, 40, 16, 16]


def generer_facture_xlsx(facture) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Facture'
    ws.sheet_view.showGridLines = False

    for i, largeur in enumerate(LARGEURS_COLONNES, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur

    if LOGO_PATH.exists():
        image = XLImage(str(LOGO_PATH))
        image.width = 85
        image.height = 75
        ws.add_image(image, 'A1')

    ws.merge_cells('C1:D1')
    titre = ws.cell(row=1, column=3, value='Facture')
    titre.font = FONT_TITRE
    titre.alignment = Alignment(horizontal='right')

    ligne = 4
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
    ws.cell(row=ligne, column=1).border = BORDURE_PRINCIPALE
    ligne += 1

    debut_bloc = ligne
    for texte, font in [
        (COMPANY_NAME_CONTACT, FONT_ENTREPRISE_NOM),
        (COMPANY_ADDRESS, FONT_MUTED),
        (f'Identifiant TVA : {COMPANY_VAT_ID}', FONT_MUTED),
        (COMPANY_EMAIL, FONT_MUTED),
    ]:
        cellule = ws.cell(row=ligne, column=1, value=texte)
        cellule.font = font
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=2)
        ligne += 1

    affaire = facture.affaire
    ligne_droite = debut_bloc
    for texte in [
        f'Date : {formater_date_longue(facture.date_facture)}',
        f'Numéro de facture : {facture.numero_facture}',
        f'Référence client : {affaire.client.matricule_fiscal or "—"}',
    ]:
        cellule = ws.cell(row=ligne_droite, column=3, value=texte)
        cellule.font = FONT_CHAMP_DROITE
        cellule.alignment = Alignment(horizontal='right')
        ws.merge_cells(start_row=ligne_droite, start_column=3, end_row=ligne_droite, end_column=4)
        ligne_droite += 1

    ligne += 1
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
    ws.cell(row=ligne, column=1).border = BORDURE_PRINCIPALE
    ligne += 1

    client = affaire.client
    cellule = ws.cell(row=ligne, column=1, value=f'À : {client.raison_sociale}')
    cellule.font = FONT_A_CLIENT
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
    ligne += 1

    coordonnees = ' · '.join(filter(None, [client.adresse, client.pays]))
    if coordonnees:
        cellule = ws.cell(row=ligne, column=1, value=coordonnees)
        cellule.font = FONT_MUTED
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
        ligne += 1

    ligne += 1

    for col, texte in enumerate(COLONNES_TABLEAU, start=1):
        cellule = ws.cell(row=ligne, column=col, value=texte)
        cellule.font = FONT_ENTETE_TABLE
        cellule.fill = FILL_ENTETE_TABLE
        cellule.border = BORDURE_GRILLE
        cellule.alignment = Alignment(horizontal='left' if col == 2 else 'right', vertical='center')
    ws.row_dimensions[ligne].height = 20
    ligne += 1

    # Première ligne du tableau (sous l'en-tête des colonnes) : le titre du
    # projet facturé, dans la même structure de colonnes que les lignes qui
    # suivent — pas une bannière fusionnée à part.
    for col in (1, 3, 4):
        cellule = ws.cell(row=ligne, column=col, value='—')
        cellule.font = FONT_NORMAL
        cellule.fill = FILL_TOTAL
        cellule.border = BORDURE_GRILLE
        cellule.alignment = Alignment(horizontal='right')
    cellule = ws.cell(row=ligne, column=2, value=affaire.devis.objet or 'Prestation')
    cellule.font = FONT_TITRE_PROJET_LIGNE
    cellule.fill = FILL_TOTAL
    cellule.border = BORDURE_GRILLE
    cellule.alignment = Alignment(horizontal='left', vertical='center')
    ligne += 1

    lignes_facture = list(facture.lignes.all())
    for i, l in enumerate(lignes_facture):
        valeurs = [
            float(l.quantite) if l.quantite is not None else None,
            l.description,
            float(l.prix_unitaire) if l.prix_unitaire is not None else None,
            float(l.montant),
        ]
        fond = FILL_ALT if i % 2 == 1 else None
        for col, valeur in enumerate(valeurs, start=1):
            cellule = ws.cell(row=ligne, column=col, value=valeur)
            cellule.font = FONT_NORMAL
            cellule.border = BORDURE_GRILLE
            if fond:
                cellule.fill = fond
            if col == 2:
                cellule.alignment = Alignment(horizontal='left')
            else:
                cellule.alignment = Alignment(horizontal='right')
                if col == 1:
                    cellule.number_format = FORMAT_NOMBRE
                elif col in (3, 4):
                    cellule.number_format = FORMAT_MONTANT
        ligne += 1

    ligne += 1
    totaux = [
        ('Sous-total', float(facture.sous_total), False),
        ('TVA', float(facture.montant_tva), False),
        ('Total', float(facture.montant_total), True),
    ]
    for label, valeur, accent in totaux:
        cellule_label = ws.cell(row=ligne, column=2, value=label)
        cellule_valeur = ws.cell(row=ligne, column=4, value=valeur)
        cellule_valeur.number_format = FORMAT_MONTANT
        cellule_valeur.alignment = Alignment(horizontal='right')
        if accent:
            cellule_label.font = FONT_TOTAL_LABEL
            cellule_valeur.font = FONT_TOTAL_VALEUR
        else:
            cellule_label.font = FONT_MUTED
            cellule_valeur.font = FONT_NORMAL
        if label == 'Total':
            cellule_label.fill = FILL_TOTAL
            cellule_valeur.fill = FILL_TOTAL
        ligne += 1

    ligne += 2
    cellule = ws.cell(row=ligne, column=1, value='Informations de paiement')
    cellule.font = Font(name='Calibri', size=12, bold=True, color=COULEUR_PRINCIPALE)
    ligne += 1
    for label, valeur in [
        ('À payer', f'{facture.montant_a_payer:,.2f} €'.replace(',', ' ').replace('.', ',')),
        ('Mode de règlement', MODE_REGLEMENT_LABELS.get(facture.mode_reglement, facture.mode_reglement)),
        ('Conditions de paiement', facture.label_echeance),
        ("Date d'échéance", formater_date_longue(facture.date_echeance)),
    ]:
        cellule_label = ws.cell(row=ligne, column=1, value=label)
        cellule_label.font = FONT_LABEL
        cellule_label.fill = FILL_ALT
        cellule_label.border = BORDURE_GRILLE
        cellule = ws.cell(row=ligne, column=2, value=str(valeur))
        cellule.font = FONT_NORMAL
        cellule.border = BORDURE_GRILLE
        ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne, end_column=4)
        for col in (3, 4):
            ws.cell(row=ligne, column=col).border = BORDURE_GRILLE
        ligne += 1

    ligne += 1
    cellule = ws.cell(row=ligne, column=1, value='Coordonnées bancaires')
    cellule.font = Font(name='Calibri', size=12, bold=True, color=COULEUR_PRINCIPALE)
    ligne += 1
    for label, valeur in [('Banque', BANQUE_NOM), ('IBAN', BANQUE_IBAN), ('BIC', BANQUE_BIC)]:
        cellule_label = ws.cell(row=ligne, column=1, value=label)
        cellule_label.font = FONT_LABEL
        cellule_label.fill = FILL_ALT
        cellule_label.border = BORDURE_GRILLE
        cellule = ws.cell(row=ligne, column=2, value=valeur)
        cellule.font = FONT_NORMAL
        cellule.border = BORDURE_GRILLE
        ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne, end_column=4)
        for col in (3, 4):
            ws.cell(row=ligne, column=col).border = BORDURE_GRILLE
        ligne += 1

    if facture.signature:
        ligne += 2
        cellule = ws.cell(row=ligne, column=1, value='Signature')
        cellule.font = Font(name='Calibri', size=12, bold=True, color=COULEUR_PRINCIPALE)
        ligne += 1
        with facture.signature.open('rb') as f:
            largeur_px, hauteur_px = PILImage.open(f).size
        largeur_max_px, hauteur_max_px = 170, 95
        ratio = min(largeur_max_px / largeur_px, hauteur_max_px / hauteur_px, 1)
        image_signature = XLImage(facture.signature.path)
        image_signature.width = largeur_px * ratio
        image_signature.height = hauteur_px * ratio
        ws.add_image(image_signature, f'A{ligne}')
        ligne += max(4, int(image_signature.height / 18) + 1)

    ligne += 2
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
    merci = ws.cell(row=ligne, column=1, value='Nous vous remercions de votre confiance !')
    merci.font = FONT_MERCI
    merci.alignment = Alignment(horizontal='center')
    ligne += 1
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
    adresse = ws.cell(row=ligne, column=1, value=f'{COMPANY_NAME_CONTACT} — {COMPANY_ADDRESS}')
    adresse.font = FONT_MUTED
    adresse.alignment = Alignment(horizontal='center')

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
